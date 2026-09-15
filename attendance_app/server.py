from fastapi import FastAPI, File, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional, List
import uvicorn
import os
import shutil
import sqlite3
import json
import hashlib
import secrets
from datetime import datetime
from deepface import DeepFace
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = FastAPI()

UPLOAD_DIR = "temp_uploads"
if not os.path.exists(UPLOAD_DIR):
    os.makedirs(UPLOAD_DIR)

MODEL_NAME = "ArcFace"
DETECTOR_BACKEND = "retinaface"
DB_PATH = r"C:\face_attendance_system\data\db"
SQLITE_DB = "attendance.db"
STUDENTS_FILE = "students.json"
USERS_FILE = "users.json"

# ========== МОДЕЛИ ==========

class StudentCreate(BaseModel):
    name: str
    group: str = ""

class StudentUpdate(BaseModel):
    name: Optional[str] = None
    group: Optional[str] = None

class UserCreate(BaseModel):
    username: str
    password: str
    role: str = "teacher"

class UserLogin(BaseModel):
    username: str
    password: str

class AttendanceSave(BaseModel):
    students: List[str]
    group: str = ""
    date: str = ""

# ========== ФУНКЦИИ ДЛЯ СТУДЕНТОВ ==========

def load_students():
    try:
        with open(STUDENTS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        default_students = [
            {"id": 1, "name": "Хмара Дмитрий Евгеньевич", "group": "22ВТ-09.03.03.01-о1"},
            {"id": 2, "name": "Петров Петр Петрович", "group": "22ВТ-09.03.03.01-о2"},
            {"id": 3, "name": "Сидорова Анна Сергеевна", "group": "22ВТ-09.03.03.01-о3"},
            {"id": 4, "name": "Козлов Дмитрий Александрович", "group": "22ВТ-09.03.03.01-о3"},
        ]
        save_students(default_students)
        return default_students

def save_students(students):
    with open(STUDENTS_FILE, "w", encoding="utf-8") as f:
        json.dump(students, f, ensure_ascii=False, indent=2)

students_db = load_students()

# ========== ФУНКЦИИ ДЛЯ ПОЛЬЗОВАТЕЛЕЙ ==========

def hash_password(password: str) -> str:
    salt = secrets.token_hex(16)
    hashed = hashlib.sha256((password + salt).encode()).hexdigest()
    return f"{salt}:{hashed}"

def verify_password(password: str, hashed: str) -> bool:
    try:
        salt, hash_value = hashed.split(":")
        return hashlib.sha256((password + salt).encode()).hexdigest() == hash_value
    except:
        return False

def load_users():
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        default_users = [
            {"id": 1, "username": "admin", "password_hash": hash_password("admin123"), "role": "admin", "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"id": 2, "username": "teacher", "password_hash": hash_password("teacher123"), "role": "teacher", "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
        ]
        save_users(default_users)
        return default_users

def save_users(users):
    with open(USERS_FILE, "w", encoding="utf-8") as f:
        json.dump(users, f, ensure_ascii=False, indent=2)

users_db = load_users()

# ========== ФУНКЦИИ БАЗЫ ДАННЫХ ==========

def init_db():
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS attendance (id INTEGER PRIMARY KEY AUTOINCREMENT, student_name TEXT NOT NULL, student_group TEXT NOT NULL, timestamp TEXT NOT NULL)''')
    conn.commit()
    conn.close()

def init_students_table():
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS students (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL UNIQUE, group_name TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL DEFAULT (datetime('now')))''')
    cursor.execute("SELECT COUNT(*) FROM students")
    if cursor.fetchone()[0] == 0:
        for student in students_db:
            try:
                cursor.execute("INSERT INTO students (name, group_name) VALUES (?, ?)", (student["name"], student.get("group", "")))
            except:
                pass
        conn.commit()
    conn.close()

def init_audit_table():
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS audit_log (id INTEGER PRIMARY KEY AUTOINCREMENT, user_name TEXT NOT NULL, action TEXT NOT NULL, details TEXT, timestamp TEXT NOT NULL DEFAULT (datetime('now')))''')
    conn.commit()
    conn.close()

def log_action(user_name: str, action: str, details: str = ""):
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute("INSERT INTO audit_log (user_name, action, details, timestamp) VALUES (?, ?, ?, ?)", (user_name, action, details, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    conn.commit()
    conn.close()

init_db()
init_students_table()
init_audit_table()

# ========== 1. РАСПОЗНАВАНИЕ ЛИЦ ==========

@app.post("/recognize")
async def recognize_faces(file: UploadFile = File(...)):
    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    try:
        results = DeepFace.find(img_path=file_path, db_path=DB_PATH, model_name=MODEL_NAME, detector_backend=DETECTOR_BACKEND, enforce_detection=False, align=True)
        present_students = set()

        for i, df in enumerate(results):
            if not df.empty:
                best_match_path = str(df.iloc[0]['identity'])
                name = Path(best_match_path).parent.name
                present_students.add(name)

        return {"status": "success", "present": list(present_students)}

    except Exception as e:
        return {"status": "error", "message": str(e)}
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)

# ========== 2. ЭКСПОРТ EXCEL ==========

@app.get("/export-excel")
async def export_excel(date: str = None, group: str = None):
    if not date:
        date = datetime.now().strftime("%Y-%m-%d")
        
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    
    if group:
        cursor.execute("SELECT student_name, student_group, timestamp FROM attendance WHERE timestamp LIKE ? AND student_group = ?", (f"{date}%", group))
    else:
        cursor.execute("SELECT student_name, student_group, timestamp FROM attendance WHERE timestamp LIKE ?", (f"{date}%",))
    
    rows = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Посещаемость"
    ws.append(["ФИО Студента", "Группа", "Время отметки"])
    
    header_font = Font(name="Calibri", bold=True, color="000000")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")

    for col in range(1, 4):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    for row_data in rows:
        ws.append([row_data[0], row_data[1], row_data[2]])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    file_path = os.path.join(UPLOAD_DIR, f"report_{date}.xlsx")
    wb.save(file_path)
    
    return FileResponse(path=file_path, filename=f"Attendance_{date}.xlsx", media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ========== 3. АНАЛИТИКА ==========

@app.get("/analytics")
async def get_analytics():
    today = datetime.now().strftime("%Y-%m-%d")
    
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    
    cursor.execute("SELECT COUNT(*) FROM attendance")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(DISTINCT student_name) FROM attendance WHERE timestamp LIKE ?", (f"{today}%",))
    today_count = cursor.fetchone()[0]
    
    cursor.execute("SELECT SUBSTR(timestamp, 1, 10) as date_day, COUNT(student_name) FROM attendance GROUP BY date_day ORDER BY date_day ASC LIMIT 5")
    chart_rows = cursor.fetchall()
    conn.close()
    
    history_data = [{"date": row[0], "count": row[1]} for row in chart_rows]
    
    return {"status": "success", "total_all_time": total_records, "total_today": today_count, "history": history_data}

# ========== 4. УПРАВЛЕНИЕ СТУДЕНТАМИ ==========

@app.get("/students")
async def get_students():
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute("SELECT id, name, group_name FROM students ORDER BY name")
    rows = cursor.fetchall()
    conn.close()
    return [{"id": row[0], "name": row[1], "group": row[2]} for row in rows]

@app.post("/students")
async def add_student(student: StudentCreate):
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO students (name, group_name) VALUES (?, ?)", (student.name, student.group))
        conn.commit()
        student_id = cursor.lastrowid
        students_db.append({"id": student_id, "name": student.name, "group": student.group})
        save_students(students_db)
        log_action("system", "Добавлен студент", f"{student.name}, группа: {student.group}")
        return {"status": "success", "student": {"id": student_id, "name": student.name, "group": student.group}}
    except sqlite3.IntegrityError:
        return {"status": "error", "message": "Студент с таким именем уже существует"}
    finally:
        conn.close()

@app.put("/students/{student_id}")
async def update_student(student_id: int, student: StudentUpdate):
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    update_fields = []
    values = []
    if student.name is not None:
        update_fields.append("name = ?")
        values.append(student.name)
    if student.group is not None:
        update_fields.append("group_name = ?")
        values.append(student.group)
    if not update_fields:
        return {"status": "error", "message": "Нет данных для обновления"}
    values.append(student_id)
    cursor.execute(f"UPDATE students SET {', '.join(update_fields)} WHERE id = ?", values)
    conn.commit()
    for s in students_db:
        if s["id"] == student_id:
            if student.name: s["name"] = student.name
            if student.group: s["group"] = student.group
    save_students(students_db)
    conn.close()
    return {"status": "success", "message": "Данные студента обновлены"}

@app.delete("/students/{student_id}")
async def delete_student(student_id: int):
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM students WHERE id = ?", (student_id,))
    row = cursor.fetchone()
    student_name = row[0] if row else "Неизвестный"
    cursor.execute("DELETE FROM students WHERE id = ?", (student_id,))
    conn.commit()
    conn.close()
    global students_db
    students_db = [s for s in students_db if s["id"] != student_id]
    save_students(students_db)
    log_action("system", "Удален студент", student_name)
    return {"status": "success", "message": "Студент удален"}

@app.get("/groups")
async def get_groups():
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT group_name FROM students WHERE group_name != '' ORDER BY group_name")
    rows = cursor.fetchall()
    conn.close()
    return ["Все группы"] + [row[0] for row in rows]

# ========== 5. АВТОРИЗАЦИЯ И ПОЛЬЗОВАТЕЛИ ==========

@app.post("/login")
async def login(user: UserLogin):
    for u in users_db:
        if u["username"] == user.username:
            if verify_password(user.password, u["password_hash"]):
                log_action(user.username, "Вход в систему")
                return {"status": "success", "user": {"id": u["id"], "username": u["username"], "role": u["role"]}}
    log_action(user.username, "Неудачная попытка входа", "Неверный пароль")
    return {"status": "error", "message": "Неверный логин или пароль"}

@app.get("/users")
async def get_users():
    return [{"id": u["id"], "username": u["username"], "role": u["role"], "created_at": u["created_at"]} for u in users_db]

@app.post("/users")
async def add_user(user: UserCreate):
    for u in users_db:
        if u["username"] == user.username:
            return {"status": "error", "message": "Пользователь уже существует"}
    new_id = max([u["id"] for u in users_db], default=0) + 1
    new_user = {"id": new_id, "username": user.username, "password_hash": hash_password(user.password), "role": user.role, "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    users_db.append(new_user)
    save_users(users_db)
    log_action("admin", "Добавлен пользователь", f"{user.username}, роль: {user.role}")
    return {"status": "success", "user": {"id": new_id, "username": user.username, "role": user.role}}

@app.delete("/users/{user_id}")
async def delete_user(user_id: int):
    global users_db
    user_to_delete = next((u for u in users_db if u["id"] == user_id), None)
    if not user_to_delete:
        return {"status": "error", "message": "Пользователь не найден"}
    if user_to_delete["role"] == "admin" and len([u for u in users_db if u["role"] == "admin"]) <= 1:
        return {"status": "error", "message": "Нельзя удалить последнего администратора"}
    users_db = [u for u in users_db if u["id"] != user_id]
    save_users(users_db)
    log_action("admin", "Удален пользователь", user_to_delete['username'])
    return {"status": "success", "message": "Пользователь удален"}

# ========== 6. ЖУРНАЛ ДЕЙСТВИЙ ==========

@app.get("/audit")
async def get_audit_log(limit: int = 50):
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    cursor.execute("SELECT user_name, action, details, timestamp FROM audit_log ORDER BY id DESC LIMIT ?", (limit,))
    rows = cursor.fetchall()
    conn.close()
    return [{"user_name": row[0], "action": row[1], "details": row[2], "timestamp": row[3]} for row in rows]

# ========== 7. СОХРАНЕНИЕ ПОСЕЩАЕМОСТИ ==========

@app.post("/save-attendance")
async def save_attendance(data: AttendanceSave):
    if not data.students:
        return {"status": "error", "message": "Список студентов пуст"}
    
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    try:
        added_count = 0
        for name in data.students:
            cursor.execute(
                "SELECT id FROM attendance WHERE student_name = ? AND timestamp LIKE ?",
                (name, f"{data.date}%")
            )
            if not cursor.fetchone():
                cursor.execute(
                    "INSERT INTO attendance (student_name, student_group, timestamp) VALUES (?, ?, ?)",
                    (name, data.group if data.group else "Общая группа", now)
                )
                added_count += 1
        
        conn.commit()
        log_action("system", "Сохранена посещаемость", f"Дата: {data.date}, добавлено: {added_count}")
        return {"status": "success", "message": f"Сохранено {added_count} студентов"}
    except Exception as e:
        conn.rollback()
        return {"status": "error", "message": str(e)}
    finally:
        conn.close()

# ========== 8. ОЧИСТКА ДАННЫХ ==========

@app.delete("/clear-attendance")
async def clear_attendance(date: str = None):
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    try:
        if date:
            cursor.execute("DELETE FROM attendance WHERE timestamp LIKE ?", (f"{date}%",))
            deleted = cursor.rowcount
            message = f"Удалено {deleted} записей за {date}"
        else:
            cursor.execute("DELETE FROM attendance")
            deleted = cursor.rowcount
            message = f"Удалены все записи ({deleted} шт.)"
        conn.commit()
        log_action("system", "Очистка посещаемости", message)
        return {"status": "success", "message": message}
    except Exception as e:
        conn.rollback()
        return {"status": "error", "message": str(e)}
    finally:
        conn.close()

@app.delete("/clear-all")
async def clear_all_data():
    conn = sqlite3.connect(SQLITE_DB)
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM attendance")
        attendance_count = cursor.rowcount
        cursor.execute("DELETE FROM students")
        students_count = cursor.rowcount
        cursor.execute("DELETE FROM audit_log")
        audit_count = cursor.rowcount
        conn.commit()
        
        global students_db
        students_db = []
        save_students(students_db)
        
        log_action("system", "Полная очистка данных", f"Удалено: посещений {attendance_count}, студентов {students_count}, логов {audit_count}")
        return {"status": "success", "message": f"Удалено: {attendance_count} посещений, {students_count} студентов, {audit_count} логов"}
    except Exception as e:
        conn.rollback()
        return {"status": "error", "message": str(e)}
    finally:
        conn.close()

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)